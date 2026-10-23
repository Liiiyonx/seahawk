from io import BytesIO
from pathlib import Path

from PIL import Image
from pptx import Presentation


SOURCE = Path(
    r"C:\Users\Liii\Desktop\科创赛道-探海灵眸-18259099729\路演PPT_探海灵眸.pptx"
)
OUTPUT = Path(r"C:\Users\Liii\Desktop\seahawk\tmp\plan_review_20261020\ppt_images")


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    presentation = Presentation(SOURCE)
    slide = presentation.slides[5]
    index = 0
    for shape in slide.shapes:
        if not hasattr(shape, "image") or shape.width != 3291840:
            continue
        index += 1
        image = shape.image
        extension = image.ext or "bin"
        destination = OUTPUT / f"device_{index}.{extension}"
        destination.write_bytes(image.blob)
        with Image.open(BytesIO(image.blob)) as opened:
            print(destination, opened.size, len(image.blob))


if __name__ == "__main__":
    main()
