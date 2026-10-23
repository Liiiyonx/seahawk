from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from pptx import Presentation
from pypdf import PdfReader


SENSITIVE_PATTERNS = [
    "李涌翔",
    "许鑫杰",
    "陈泽韩",
    "李泽群",
    "晏可鑫",
    "范忆梅",
    "福州理工学院",
    "福州理工",
    "计算与信息科学学院",
    "智能科学与技术",
    "电子信息工程",
    "计算机科学与技术",
    "软件工程",
    "指导教师",
    "队长",
    "18259099729",
    "18259083908",
    "18059908021",
    "17850629206",
    "18965625753",
    "15170331887",
    "@qq.com",
]


def docx_text(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    return "\n".join(parts)


def pptx_text(path: Path) -> str:
    presentation = Presentation(path)
    parts: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        parts.append(f"\n--- SLIDE {slide_number} ---")
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                parts.append(shape.text)
            if hasattr(shape, "table"):
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(f"[NOTES] {notes}")
    return "\n".join(parts)


def pdf_text(path: Path) -> str:
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return docx_text(path)
    if suffix == ".pptx":
        return pptx_text(path)
    if suffix == ".pdf":
        return pdf_text(path)
    if suffix in {".txt", ".md", ".json", ".yaml", ".yml", ".py", ".js", ".vue"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    return ""


def summarize(path: Path, text: str, quiet: bool = False) -> None:
    print(f"\n===== {path} =====")
    print(f"CHARS={len(text)}")
    findings: list[tuple[str, int]] = []
    for pattern in SENSITIVE_PATTERNS:
        count = text.count(pattern)
        if count:
            findings.append((pattern, count))
    if findings:
        print("SENSITIVE=" + "; ".join(f"{value}:{count}" for value, count in findings))
    else:
        print("SENSITIVE=无命中")

    if path.suffix.lower() == ".pptx":
        presentation = Presentation(path)
        print(f"SLIDES={len(presentation.slides)}")

    if not quiet:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for line in lines[:40]:
            print(line[:500])


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("paths", nargs="+")
    parser.add_argument("--full", action="store_true")
    parser.add_argument("--quiet", action="store_true")
    args = parser.parse_args()

    for raw_path in args.paths:
        path = Path(raw_path)
        text = extract(path)
        if args.full:
            print(f"\n===== {path} =====")
            print(text)
        else:
            summarize(path, text, quiet=args.quiet)


if __name__ == "__main__":
    main()
