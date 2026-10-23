from __future__ import annotations

import re
import zipfile
from pathlib import Path

from docx import Document
from PIL import ExifTags, Image
from pptx import Presentation
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parent / "submission_build"

SENSITIVE_TERMS = (
    "李涌翔",
    "许鑫杰",
    "陈泽韩",
    "李泽群",
    "晏可鑫",
    "范忆梅",
    "福州理工学院",
    "福州理工",
    "计算与信息科学学院",
    "智能科学与技术",
    "电子信息工程",
    "计算机科学与技术",
    "软件工程",
    "指导教师",
    "队长",
    "18259099729",
    "18259083908",
    "18059908021",
    "17850629206",
    "18965625753",
    "15170331887",
)

SCAN_PATTERNS = {
    "个人目录路径": re.compile(
        r"(?:[A-Za-z]:\\{1,2}(?:Users|Documents and Settings)\\|"
        r"/mnt/[A-Za-z]/Users/|/Users/[^/\s]+/|D:\\{1,2}seahawk-work)"
    ),
    "手机号": re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)"),
    "邮箱": re.compile(
        r"(?<![\w.+-])[\w.+-]+@"
        r"(?!(?:pytest|asyncio|app|router|staticmethod|classmethod|property|wraps|fixture|mark|abstractmethod)\.)"
        r"[\w.-]+\.[A-Za-z]{2,}"
    ),
    "身份证号": re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)"),
    "疑似密钥": re.compile(
        r"(?:sk-[A-Za-z0-9_-]{20,}|AKIA[0-9A-Z]{16}|"
        r"gh[pousr]_[A-Za-z0-9]{20,}|-----BEGIN [A-Z ]*PRIVATE KEY-----)"
    ),
}

OFFICE_SUFFIXES = {".docx", ".pptx", ".xlsx"}
TEXT_SUFFIXES = {
    ".md",
    ".txt",
    ".json",
    ".yaml",
    ".yml",
    ".py",
    ".js",
    ".mjs",
    ".ts",
    ".tsx",
    ".jsx",
    ".vue",
    ".html",
    ".css",
    ".sql",
    ".toml",
    ".ini",
    ".conf",
    ".example",
    ".env",
    ".gitignore",
    ".gitattributes",
}


def clip(text: str, limit: int = 180) -> str:
    return " ".join(text.split())[:limit]


def findings_for(text: str) -> list[tuple[str, str]]:
    findings: list[tuple[str, str]] = []
    for term in SENSITIVE_TERMS:
        if term in text:
            findings.append((f"敏感词:{term}", clip(text)))
    for label, pattern in SCAN_PATTERNS.items():
        match = pattern.search(text)
        if match:
            findings.append((label, clip(text[max(0, match.start() - 40) : match.end() + 40])))
    return findings


def extract_docx(path: Path) -> str:
    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    for section in document.sections:
        parts.extend(paragraph.text for paragraph in section.header.paragraphs)
        parts.extend(paragraph.text for paragraph in section.footer.paragraphs)
    props = document.core_properties
    parts.extend(
        str(value)
        for value in (
            props.author,
            props.last_modified_by,
            props.title,
            props.subject,
            props.keywords,
            props.comments,
            props.category,
            props.identifier,
        )
        if value
    )
    return "\n".join(parts)


def extract_pptx(path: Path) -> str:
    presentation = Presentation(path)
    parts: list[str] = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                parts.append(shape.text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    parts.extend(cell.text for cell in row.cells)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            parts.append(slide.notes_slide.notes_text_frame.text)
    props = presentation.core_properties
    parts.extend(
        str(value)
        for value in (
            props.author,
            props.last_modified_by,
            props.title,
            props.subject,
            props.keywords,
            props.comments,
            props.category,
            props.identifier,
        )
        if value
    )
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    reader = PdfReader(path)
    parts = [page.extract_text() or "" for page in reader.pages]
    for key, value in (reader.metadata or {}).items():
        parts.append(f"{key}={value}")
    return "\n".join(parts)


def scan_office_parts(path: Path) -> list[tuple[str, str, str]]:
    findings: list[tuple[str, str, str]] = []
    with zipfile.ZipFile(path) as archive:
        for entry in archive.infolist():
            if not entry.filename.lower().endswith((".xml", ".rels")):
                continue
            text = archive.read(entry).decode("utf-8", errors="ignore")
            for label, evidence in findings_for(text):
                findings.append((str(path), f"OOXML:{entry.filename}:{label}", evidence))
    return findings


def scan_image_exif(path: Path) -> list[tuple[str, str, str]]:
    findings: list[tuple[str, str, str]] = []
    with Image.open(path) as image:
        exif = image.getexif()
        for tag_id, value in exif.items():
            tag_name = ExifTags.TAGS.get(tag_id, str(tag_id))
            text = f"{tag_name}={value}"
            for label, evidence in findings_for(text):
                findings.append((str(path), f"EXIF:{label}", evidence))
    return findings


def main() -> None:
    all_findings: list[tuple[str, str, str]] = []
    summary_errors: list[str] = []
    files = sorted(path for path in ROOT.rglob("*") if path.is_file())

    for path in files:
        relative = path.relative_to(ROOT)
        for label, evidence in findings_for(str(relative)):
            all_findings.append((str(relative), f"文件名:{label}", evidence))

        suffix = path.suffix.lower()
        try:
            if suffix == ".docx":
                text = extract_docx(path)
            elif suffix == ".pptx":
                text = extract_pptx(path)
            elif suffix == ".pdf":
                text = extract_pdf(path)
            elif suffix in TEXT_SUFFIXES or path.name in {".env.example", "Makefile"}:
                text = path.read_text(encoding="utf-8", errors="ignore")
            else:
                text = ""
            for label, evidence in findings_for(text):
                all_findings.append((str(relative), f"内容:{label}", evidence))
        except Exception as exc:
            summary_errors.append(f"{relative}: {type(exc).__name__}: {exc}")

        if suffix in OFFICE_SUFFIXES:
            try:
                all_findings.extend(scan_office_parts(path))
            except Exception as exc:
                summary_errors.append(f"{relative}: OOXML scan failed: {exc}")

        if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            try:
                all_findings.extend(scan_image_exif(path))
            except Exception as exc:
                summary_errors.append(f"{relative}: EXIF scan failed: {exc}")

    total_bytes = sum(path.stat().st_size for path in files)
    print(f"FILES={len(files)}")
    print(f"SIZE_MB={total_bytes / 1024 / 1024:.2f}")
    print(f"FINDINGS={len(all_findings)}")
    for path, label, evidence in all_findings:
        print(f"HIT | {path} | {label} | {evidence}")
    print(f"SCAN_ERRORS={len(summary_errors)}")
    for error in summary_errors:
        print(f"ERROR | {error}")


if __name__ == "__main__":
    main()
