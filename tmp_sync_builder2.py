#!/usr/bin/env python3
"""Sync award-polish content into build_huawei_ict_final_docx.py CONTENT."""
from pathlib import Path
import ast

p = Path(r"C:\Users\Liii\Desktop\seahawk\scripts\build_huawei_ict_final_docx.py")
s = p.read_text(encoding="utf-8")

# --- 1) insert_table centering ---
old = '    table.style = "Table Grid"'
assert old in s
s = s.replace(old, old + '\n    table.alignment = WD_TABLE_ALIGNMENT.CENTER', 1)
if "from docx.enum.table import WD_TABLE_ALIGNMENT" not in s:
    s = s.replace("from docx.enum.text import WD_ALIGN_PARAGRAPH",
                  "from docx.enum.table import WD_TABLE_ALIGNMENT\nfrom docx.enum.text import WD_ALIGN_PARAGRAPH")

# --- 2) competitor comparison after the 可进化 paragraph (line anchor) ---
anchor1 = '("p", "“可进化”落在三个层面。'
idx = s.index(anchor1)
line_end = s.index("\n", idx) + 1
competitor = (
    '        ("p", "与国内外同类方案对比：市面海漂垃圾治理方案以单机清洁机器人产品为主（WasteShark、Clearbot 等），本项目把战场从“看得见、捞得着”推进到“可追责、可复核”的治理决策层——增量在决策与数据资产层，不在单机感知硬件："),\n'
    '        ("table",\n'
    '         ["对比维度", "国内外同类方案（WasteShark / Clearbot 等）", "探海灵眸 Oceanus"],\n'
    '         [\n'
    '             ["形态定位", "单机清洁机器人产品，能力封闭在机身", "可审计治理智能体；感知—决策—执行三级载体，执行端可插拔"],\n'
    '             ["决策方式", "人工遥控或厂商内置固定规划", "智能体规划 → 策略守卫 → 多角色人工审批 → 审计回放"],\n'
    '             ["新垃圾类别接入", "需厂商改产品并重新训练", "开放词汇零样本通道，现场加词，不重训不重启"],\n'
    '             ["治理数据", "清理量等运营记录", "六类资产版本化 + 动态本体 + 决策证据链，可按版本复核"],\n'
    '             ["成本路线", "WasteShark 起售价约 2.36 万美元（厂商公开报价，E1 检索）", "国产化软件定义路线，执行端按需选型、不被单机厂商绑定"],\n'
    '         ]),\n'
)
s = s[:line_end] + competitor + s[line_end:]

# --- 3) business value section before 诚实边界 + renumber ---
old_h3 = '        ("h3", "六、诚实边界声明"),'
assert old_h3 in s
new_block = (
    '        ("h3", "六、业务价值测算与推广前景"),\n'
    '        ("p", "治理链路四个环节的效率对照（系统侧为 E1/E2 实测，人工基线为 E0 假设口径；测算口径，非真实运营收益）："),\n'
    '        ("table",\n'
    '         ["环节", "人工基线（E0 假设）", "系统链路（E1/E2 实测）"],\n'
    '         [\n'
    '             ["跨文档政策/台账检索", "跨系统人工查找，10–30 分钟/次", "知识域多跳检索 <2 秒，返回带引用证据链（R-KN-01 / R-KN-04）"],\n'
    '             ["决策依据整理", "人工整理证据清单，15–30 分钟/份", "决策自动固化 4 条有序证据，<5 秒（R-KN-04）"],\n'
    '             ["事件派单", "电话/口头派单，5–10 分钟", "MQTT 秒级下发，回执闭环（E2 联调实测）"],\n'
    '             ["本体维护", "人工梳理，小时级", "增量抽取较全量重抽节省 63.19% 耗时（R-KN-03）"],\n'
    '         ]),\n'
    '        ("p", "多行业复制路径已有演示级验证：同一套 5 个 Skill 工作流模板在海洋治理、医疗、政务三领域语料上完成轻量化迁移（R-MG-01，检索命中 3/3），迁移仅替换资产源、领域词表、标准号、问题集与角色映射五个变更点；该验证为 E1 演示级口径，非真实行业交付。"),\n'
    '        ("p", "知识产权：软件著作权《探海灵眸海洋环境治理智能体软件 V1.0》（著作权人：福州理工学院）登记申请材料已按中国版权保护中心要求制备完成（60 页源代码鉴别材料、软件说明书、约 9 万行源程序量统计），单位申请流程进行中。政策背景：《中华人民共和国海洋环境保护法》（2023 年修订）强化海洋垃圾治理要求，沿海县域治理需求走向制度化、常态化，为同类场景复制提供需求基础。"),\n'
    '        ("h3", "七、诚实边界声明"),\n'
)
s = s.replace(old_h3, new_block)

# --- 4) footer page numbers in main(): find document.save call ---
save_anchor = None
for cand in ("    document.save(", "    output_path.parent.mkdir"):
    if cand in s:
        save_anchor = cand
        break
assert save_anchor, "save anchor not found"
footer_fn = '''
def add_footer_page_numbers(doc) -> None:
    """Centered '第 X 页 / 共 Y 页' in footer; cover kept clean via titlePg."""
    from docx.shared import RGBColor as _RGB
    sec = doc.sections[0]
    sec.different_first_page_header_footer = True
    fp = sec.footer.paragraphs[0]
    fp.text = ""
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    gray = _RGB(0x59, 0x59, 0x59)

    def _field(par, instr):
        run = par.add_run()
        b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
        i = OxmlElement("w:instrText"); i.set(qn("xml:space"), "preserve"); i.text = instr
        sp = OxmlElement("w:fldChar"); sp.set(qn("w:fldCharType"), "separate")
        t = OxmlElement("w:t"); t.text = "1"
        e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
        run._r.extend([b, i, sp, t, e])
        run.font.size = Pt(9)
        run.font.color.rgb = gray

    def _txt(par, text):
        r = par.add_run(text)
        r.font.size = Pt(9)
        r.font.color.rgb = gray

    _txt(fp, "第 ")
    _field(fp, " PAGE ")
    _txt(fp, " 页 / 共 ")
    _field(fp, " NUMPAGES ")
    _txt(fp, " 页")

'''
s = s.replace("def main() -> None:", footer_fn + "\ndef main() -> None:", 1)
s = s.replace(save_anchor, "    add_footer_page_numbers(doc)\n" + save_anchor, 1)

p.write_text(s, encoding="utf-8")
ast.parse(s)
print("builder synced, syntax OK; save anchor was:", save_anchor)
