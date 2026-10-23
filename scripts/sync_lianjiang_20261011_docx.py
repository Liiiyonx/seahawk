# -*- coding: utf-8 -*-
"""连江材料同步 · 创业项目计划书 docx 更新（2026-10-11）。

与 xlsx 同一口径：27/27 场景、SC-FloatWaste 公开基准、42%~76% 误报区间、
127 项故障注入、13 家 15 份询价、海渔局访谈、软著登记申请中、SeaSight→Oceanus。
"""
from __future__ import annotations

import copy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

DOCX = Path(r"C:\Users\Liii\Desktop\连江县创新创业大赛\创业项目计划书_探海灵眸_初稿.docx")


def rewrite_para(p, text: str) -> None:
    """整段重写：保留首 run 的格式（rPr），替换全部文本。"""
    runs = p.runs
    src = runs[0] if runs else None
    rpr = src._element.find(qn("w:rPr")) if src is not None else None
    for r in list(runs):
        r._element.getparent().remove(r._element)
    new_run = p.add_run(text)
    if rpr is not None:
        new_run._element.insert(0, copy.deepcopy(rpr))


def rewrite_cell(cell, text: str) -> None:
    """单元格重写：保留首段首 run 的格式。"""
    p = cell.paragraphs[0]
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)
    rewrite_para(p, text)


def append_to_para(p, suffix: str) -> None:
    rewrite_para(p, p.text + suffix)


def replace_in_runs(p, old: str, new: str) -> None:
    """仅替换 run 文本中的子串（结构简单时用）。"""
    for r in p.runs:
        if old in r.text:
            r.text = r.text.replace(old, new)


PARA_46 = (
    "项目已具备完整软件侧证据：1,061 项自动化测试全部通过（0 失败/0 跳过），"
    "Agent 固定场景评测 27/27，真实状态软件集成验收 129/129，受控故障注入自证 127 项通过。"
    "视觉侧采用零标注数据依赖的 OpenCV 时序校验方案，在嵌入式纯 CPU 上完成推理，"
    "时序误报抑制率 42%～76%；并完成第三方公开漂浮垃圾基准（SC-FloatWaste，400 图/1,490 框）"
    "独立评测：监督微调通道 mAP50 0.514、零样本开放词汇通道 mAP50 0.116"
    "（口径：公开基准独立评测精度，非连江现场验证）。感知精度按 not_evaluated 如实标注，"
    "不以链路跑通代替识别精度。机械臂执行端已在受控仿真海面完成“派单—拾取—回传”闭环；"
    "供应链侧完成 5 类核心部件、13 家供应商、15 份询价（含 2 份盖章报价单），"
    "用户侧完成连江县海渔局访谈（记录存档）；软件著作权登记申请中。"
)

PARA_55_SUFFIX = (
    "平台已完成演示环境搭建——含监测大屏、事件中心、工单看板、设备态势、治理报表、"
    "智能助手、智能体控制台、知识智能体、双通道对照、机械臂执行端等 10 个页面，"
    "支持四类角色分级权限；软件著作权登记申请中。"
)

PARA_101 = (
    "近期（赛事前）：已完成连江县海渔局访谈与 5 类核心部件供应商询价（13 家/15 份），"
    "补齐受控验证证据并完成演示录屏；继续推进机械臂真机对接与乡镇试点意向。"
)

PARA_108 = (
    "行动卡：连江县海渔局访谈与核心部件供应商询价已完成；继续推进乡镇海渔站/环卫队"
    "访谈预约与试点意向落地。"
)

NEW_ROWS_T0 = [
    ("公开基准评测", "第三方公开漂浮垃圾基准 SC-FloatWaste（400 图 / 1,490 框）独立评测",
     "E1；监督微调 mAP50 0.514、零样本 0.116（非现场验证）"),
    ("用户与供应链证据", "连江县海渔局访谈完成（记录存档）；5 类核心部件 13 家供应商 15 份询价（含 2 份盖章报价单）",
     "原始记录可溯源，可复核"),
]


def main() -> None:
    doc = Document(DOCX)

    # ---- 正文段落 ----
    for idx in (5, 8, 44):
        replace_in_runs(doc.paragraphs[idx], "SeaSight", "Oceanus")
    rewrite_para(doc.paragraphs[46], PARA_46)
    append_to_para(doc.paragraphs[55], PARA_55_SUFFIX)
    rewrite_para(doc.paragraphs[101], PARA_101)
    rewrite_para(doc.paragraphs[108], PARA_108)

    # ---- 表0 能力表 ----
    t0 = doc.tables[0]
    rewrite_cell(t0.rows[2].cells[2], "E1；27/27 固定场景通过")
    rewrite_cell(t0.rows[3].cells[2], "E2；129/129 通过；受控故障注入 127 项")
    rewrite_cell(t0.rows[5].cells[1], "1,061 项自动化测试 0 失败/0 跳过，契约漂移与受控故障注入自证 127 项")
    # 追加两行（复制末行 XML 保格式）
    tbl = t0._tbl
    for cells in NEW_ROWS_T0:
        new_tr = copy.deepcopy(t0.rows[-1]._tr)
        tbl.append(new_tr)
        row = t0.rows[-1]
        for ci, text in enumerate(cells):
            rewrite_cell(row.cells[ci], text)

    # ---- 表4 风险表 ----
    t4 = doc.tables[4]
    rewrite_cell(t4.rows[1].cells[3], "访谈与询价已完成；持续补受控/现场验证证据；不能取得的证据主动标注")
    rewrite_cell(t4.rows[2].cells[3], "坚持 not_evaluated 口径，不以 42%～76% 代替识别精度")

    # ---- 页眉 ----
    for sec in doc.sections:
        for p in sec.header.paragraphs:
            replace_in_runs(p, "SeaSight", "Oceanus")

    doc.save(DOCX)
    print("[OK] 已更新:", DOCX.name)


if __name__ == "__main__":
    main()