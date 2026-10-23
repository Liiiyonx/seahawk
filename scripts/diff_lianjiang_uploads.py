# -*- coding: utf-8 -*-
"""对比用户上传的4个连江材料 vs 桌面版是否一致"""
from docx import Document
from pptx import Presentation
from pypdf import PdfReader
import openpyxl

PAIRS = [
    ("xlsx", r"c:\Users\Liii\.trae-cn\attachments\6ac9d744f7615d7951be5ad4\e18d4d38-919d-4a46-80b0-af70ef576ea5_9e2258e3-b420-47c7-8600-0050e9f94080_报名表_探海灵眸_初稿.xlsx",
     r"C:\Users\Liii\Desktop\连江县创新创业大赛\报名表_探海灵眸_初稿.xlsx"),
    ("docx", r"c:\Users\Liii\.trae-cn\attachments\6ac9d744f7615d7951be5ad4\bbd77823-2ccb-4190-9543-24e9b617f6de_c9f60023-b8ed-491f-ba26-978a1559c46d_创业项目计划书_探海灵眸_初稿.docx",
     r"C:\Users\Liii\Desktop\连江县创新创业大赛\创业项目计划书_探海灵眸_初稿.docx"),
    ("pdf", r"c:\Users\Liii\.trae-cn\attachments\6ac9d744f7615d7951be5ad4\0d227435-fa73-411d-92bf-7bcff85ffe57_26dfebc0-113c-486a-be27-b074b2cc94f8_创业项目计划书_探海灵眸_排版预览.pdf",
     r"C:\Users\Liii\Desktop\连江县创新创业大赛\创业项目计划书_探海灵眸_排版预览.pdf"),
    ("pptx", r"c:\Users\Liii\.trae-cn\attachments\6ac9d744f7615d7951be5ad4\42f5f451-8af0-4a5d-9491-796167583de6_391745a0-ef12-4cbb-8579-d52f7e6491b1_探海灵眸Oceanus_路演PPT_潮汐守望队.pptx",
     r"C:\Users\Liii\Desktop\连江县创新创业大赛\探海灵眸Oceanus_路演PPT_潮汐守望队.pptx"),
]


def read_xlsx(p):
    wb = openpyxl.load_workbook(p, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append("## " + ws.title)
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None:
                    out.append("%s=%s" % (c.coordinate, c.value))
    return out


def read_docx(p):
    d = Document(p)
    out = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            out.append("|".join(c.text for c in row.cells))
    return out


def read_pdf(p):
    return [pg.extract_text() or "" for pg in PdfReader(p).pages]


def read_pptx(p):
    pr = Presentation(p)
    out = []
    for i, slide in enumerate(pr.slides):
        for sh in slide.shapes:
            if sh.has_text_frame:
                out.append("s%d:%s" % (i + 1, sh.text_frame.text[:120]))
    return out


READERS = {"xlsx": read_xlsx, "docx": read_docx, "pdf": read_pdf, "pptx": read_pptx}

for kind, up, desk in PAIRS:
    a = READERS[kind](up)
    b = READERS[kind](desk)
    if a == b:
        print("[一致] %s" % up.split("_")[-1])
    else:
        print("[差异] %s : 上传%d项 vs 桌面%d项" % (up.split("_")[-1], len(a), len(b)))
        sa, sb = set(a), set(b)
        only_a = [x for x in a if x not in sb][:8]
        only_b = [x for x in b if x not in sa][:8]
        for x in only_a:
            print("  仅上传版:", x[:100])
        for x in only_b:
            print("  仅桌面版:", x[:100])