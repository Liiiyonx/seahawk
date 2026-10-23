import sys


def inspect_xlsx(path):
    import openpyxl

    wb = openpyxl.load_workbook(path)
    print("SHEETS", wb.sheetnames)
    for ws in wb.worksheets:
        print("=" * 80)
        print("SHEET", ws.title, ws.max_row, ws.max_column)
        for row in ws.iter_rows():
            for c in row:
                if c.value is None:
                    continue
                s = str(c.value)
                if any(k in s for k in ["\u8f6f\u8457", "\u8f6f\u4ef6\u8457\u4f5c\u6743", "\u77e5\u8bc6\u4ea7\u6743", "\u767b\u8bb0"]):
                    print(repr(c.coordinate), repr(s))


def inspect_docx(path):
    from docx import Document

    d = Document(path)
    print("PARAS", len(d.paragraphs), "TABLES", len(d.tables))
    for i, para in enumerate(d.paragraphs):
        t = para.text.strip()
        if t:
            print(i, "|", para.style.name, "|", t)
    print("=" * 40)
    for ti, tb in enumerate(d.tables):
        print("TABLE", ti, len(tb.rows), len(tb.columns))
        for r in tb.rows:
            print(" | ".join(c.text.strip().replace("\n", " / ") for c in r.cells))


def inspect_pdf(path):
    import fitz

    doc = fitz.open(path)
    print("PAGES", doc.page_count)
    for i, page in enumerate(doc):
        t = page.get_text()
        if "\u8f6f\u4ef6\u8457\u4f5c\u6743" in t or "\u8f6f\u8457" in t:
            print("=" * 80)
            print("PAGE", i + 1)
            print(t)


if __name__ == "__main__":
    kind, path = sys.argv[1], sys.argv[2]
    {"xlsx": inspect_xlsx, "docx": inspect_docx, "pdf": inspect_pdf}[kind](path)
