#!/usr/bin/env python3
"""Font & layout audit for the two submission docx files."""
from __future__ import annotations

import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

PKG = Path(r"C:\Users\Liii\Desktop\华为ICT大赛\提交包")
FILES = [
    ("材料1", PKG / "01_开发设计文档.docx"),
    ("材料2", PKG / "02_智能体设计说明.docx"),
]

EMU_PER_CM = 360000


def cm(v):
    return round(v / EMU_PER_CM, 2) if v is not None else None


def run_font_info(run):
    rPr = run._element.rPr
    ascii_f = ea_f = None
    if rPr is not None and rPr.rFonts is not None:
        ascii_f = rPr.rFonts.get(qn("w:ascii"))
        ea_f = rPr.rFonts.get(qn("w:eastAsia"))
    size = run.font.size.pt if run.font.size else None
    return ascii_f, ea_f, size


def audit(tag, path):
    doc = Document(path)
    print(f"\n{'=' * 70}\n{tag}: {path.name}")
    sec = doc.sections[0]
    print(f"页面: {cm(sec.page_width)}cm x {cm(sec.page_height)}cm  "
          f"边距 上{cm(sec.top_margin)}/下{cm(sec.bottom_margin)}/"
          f"左{cm(sec.left_margin)}/右{cm(sec.right_margin)}cm  "
          f"可用宽 {cm(sec.page_width - sec.left_margin - sec.right_margin)}cm")
    usable = sec.page_width - sec.left_margin - sec.right_margin

    normal = doc.styles["Normal"]
    nf = normal.font
    ea = None
    if normal._element.rPr is not None and normal._element.rPr.rFonts is not None:
        ea = normal._element.rPr.rFonts.get(qn("w:eastAsia"))
    print(f"Normal 样式: {nf.name} / eastAsia={ea} / {nf.size.pt if nf.size else None}pt")

    # heading styles
    for sname in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[sname]
            f = st.font
            hea = None
            if st._element.rPr is not None and st._element.rPr.rFonts is not None:
                hea = st._element.rPr.rFonts.get(qn("w:eastAsia"))
            print(f"  {sname}: {f.name} / eastAsia={hea} / {f.size.pt if f.size else None}pt")
        except KeyError:
            print(f"  {sname}: <无>")

    # run-level font census
    font_counter = Counter()
    ea_missing = 0
    tiny = []
    huge = []
    total_runs = 0

    def scan_runs(paras, where):
        nonlocal ea_missing, total_runs
        for p in paras:
            for r in p.runs:
                if not r.text.strip():
                    continue
                total_runs += 1
                a, e, s = run_font_info(r)
                font_counter[(a, e, s)] += 1
                if e is None and a is None:
                    ea_missing += 1
                if s is not None and s < 7:
                    tiny.append((where, s, r.text[:18]))
                if s is not None and s > 20 and tag == "材料2":
                    huge.append((where, s, r.text[:18]))

    scan_runs(doc.paragraphs, "body")
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                scan_runs(c.paragraphs, "table")

    print(f"正文 run 总数: {total_runs}  无任何字体声明: {ea_missing}")
    print("字体组合 Top8 (ascii, eastAsia, size_pt -> 出现次数):")
    for (a, e, s), n in font_counter.most_common(8):
        print(f"   {a} | {e} | {s}pt -> {n}")
    if tiny:
        print(f"  ⚠ <7pt 小字 run {len(tiny)} 处: {tiny[:5]}")
    if huge:
        print(f"  ⚠ >20pt run: {huge[:3]}")

    # eastAsia-only check: runs whose ascii font is CJK-incompatible default
    odd_fonts = Counter()
    for (a, e, s), n in font_counter.items():
        if a not in ("Microsoft YaHei", "Consolas", "微软雅黑", None):
            odd_fonts[(a, e)] += n
    if odd_fonts:
        print("  ⚠ 非常规 ascii 字体:", dict(odd_fonts))

    # tables: width vs usable
    over = 0
    for i, t in enumerate(doc.tables):
        try:
            row = t.rows[0]
            widths = []
            seen = set()
            for c in row.cells:
                if id(c._tc) in seen:
                    continue
                seen.add(id(c._tc))
                widths.append(c.width)
            total_w = sum(w for w in widths if w)
            if total_w and total_w > usable + Cm(0.05):
                over += 1
                print(f"  ⚠ 表{i}: 总宽 {cm(total_w)}cm 超过可用宽 {cm(usable)}cm")
        except Exception:
            pass
    if not over:
        print(f"表格: {len(doc.tables)} 张, 未检出超宽")

    # images
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
          "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
          "a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
    imgs = doc.element.body.findall(".//w:drawing", ns)
    widths_cm = []
    for d in imgs:
        ext = d.find(".//wp:extent", ns)
        if ext is not None:
            widths_cm.append(round(int(ext.get("cx")) / EMU_PER_CM, 2))
    print(f"图片: {len(imgs)} 张, 宽度cm = {sorted(set(widths_cm), reverse=True)}")
    too_wide = [w for w in widths_cm if w > cm(usable) + 0.05]
    if too_wide:
        print(f"  ⚠ 超可用宽的图片: {too_wide}")

    # heading page breaks & spacing anomalies (材料2 built with page_break_before)
    h1s = [p for p in doc.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    print(f"H1 数量: {len(h1s)}")

    # paragraph line spacing census (body only)
    ls = Counter()
    for p in doc.paragraphs:
        if p.text.strip():
            v = p.paragraph_format.line_spacing
            ls[str(v)] += 1
    print("行距分布:", dict(ls.most_common(5)))


for tag, p in FILES:
    audit(tag, p)
print("\nDONE")
