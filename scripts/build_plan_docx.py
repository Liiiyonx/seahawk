#!/usr/bin/env python3
"""Build the Oceanus proposal DOCX from the controlled Markdown source."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLACK = "000000"
HEADER_FILL = "1F4E79"
ALT_FILL = "F3F7FA"
BORDER_COLOR = "D9D9D9"
CODE_FILL = "F5F7F9"
BODY_FONT = "Microsoft YaHei"
BODY_EAST_ASIA_FONT = "Microsoft YaHei"
MONO_FONT = "Consolas"


@dataclass
class Block:
    kind: str
    text: str = ""
    level: int = 0
    rows: list[list[str]] | None = None
    lines: list[str] | None = None
    ordered: bool = False
    marker: str = ""


def _normalize_heading(text: str) -> str:
    return re.sub(r"\s+", "", text).strip()


def parse_markdown(source: str) -> tuple[dict[str, str], list[Block]]:
    lines = source.splitlines()
    title = ""
    metadata: dict[str, str] = {}
    body: list[str] = []
    in_cover = True

    for line in lines:
        if in_cover and line.startswith("# "):
            title = line[2:].strip()
            continue
        if in_cover:
            match = re.match(r">\s*([^：]+)：(.*)$", line)
            if match:
                metadata[match.group(1).strip()] = match.group(2).strip()
                continue
            if line.strip():
                in_cover = False
        body.append(line)

    blocks: list[Block] = []
    index = 0
    while index < len(body):
        raw = body[index]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            index += 1
            continue

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            code_lines: list[str] = []
            index += 1
            while index < len(body) and not body[index].strip().startswith("```"):
                code_lines.append(body[index].rstrip())
                index += 1
            if index < len(body):
                index += 1
            blocks.append(Block(kind="code", text=language, lines=code_lines))
            continue

        if stripped.startswith("|"):
            table_lines: list[str] = []
            while index < len(body) and body[index].strip().startswith("|"):
                table_lines.append(body[index].strip())
                index += 1
            rows: list[list[str]] = []
            for row_index, table_line in enumerate(table_lines):
                cells = [cell.strip() for cell in table_line.strip("|").split("|")]
                if row_index == 1 and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    continue
                rows.append(cells)
            if rows:
                blocks.append(Block(kind="table", rows=rows))
            continue

        heading = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading:
            # The controlled source uses the document title as H1; numbered
            # major sections start at H2 and map to Word Heading 1.
            source_level = len(heading.group(1))
            blocks.append(
                Block(
                    kind="heading",
                    level=max(1, source_level - 1),
                    text=heading.group(2).strip(),
                )
            )
            index += 1
            continue

        if re.fullmatch(r"-{3,}", stripped):
            blocks.append(Block(kind="spacer"))
            index += 1
            continue

        quote = re.match(r"^>\s?(.*)$", stripped)
        if quote:
            blocks.append(Block(kind="quote", text=quote.group(1).strip()))
            index += 1
            continue

        list_item = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if list_item:
            indent = len(list_item.group(1).replace("\t", "    "))
            marker = list_item.group(2)
            blocks.append(
                Block(
                    kind="list",
                    level=min(indent // 2, 3),
                    text=list_item.group(3).strip(),
                    ordered=marker.endswith("."),
                    marker=marker,
                )
            )
            index += 1
            continue

        paragraph_lines = [stripped]
        index += 1
        while index < len(body):
            next_line = body[index].strip()
            if (
                not next_line
                or next_line.startswith("#")
                or next_line.startswith("```")
                or next_line.startswith("|")
                or next_line.startswith(">")
                or re.match(r"^([-*]|\d+\.)\s+", next_line)
                or re.fullmatch(r"-{3,}", next_line)
            ):
                break
            paragraph_lines.append(next_line)
            index += 1
        blocks.append(Block(kind="paragraph", text=" ".join(paragraph_lines)))

    return {"title": title, **metadata}, blocks


def set_run_font(run, name: str, size: float | None = None, bold: bool | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_cell_margins(cell, top: int = 60, start: int = 80, bottom: int = 60, end: int = 80) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = BORDER_COLOR, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_paragraph_shading(paragraph, fill: str) -> None:
    paragraph_format = paragraph._p.get_or_add_pPr()
    shading = paragraph_format.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        paragraph_format.append(shading)
    shading.set(qn("w:fill"), fill)


def remove_paragraph_border(paragraph_or_style) -> None:
    paragraph = getattr(paragraph_or_style, "_p", paragraph_or_style)
    paragraph_format = paragraph.get_or_add_pPr()
    border = paragraph_format.find(qn("w:pBdr"))
    if border is not None:
        paragraph_format.remove(border)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, BODY_FONT, 8.5)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.20

    for style_name, size, before, after in (
        ("Heading 1", 18, 0, 11),
        ("Heading 2", 14, 14, 7),
        ("Heading 3", 11.5, 10, 5),
    ):
        style = styles[style_name]
        style.font.name = BODY_FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    title = styles["Title"]
    title.font.name = BODY_FONT
    title._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
    title.font.size = Pt(30)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    remove_paragraph_border(title.element)

    for list_style in ("List Bullet", "List Number"):
        style = styles[list_style]
        style.font.name = BODY_FONT
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), BODY_EAST_ASIA_FONT)
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(2)
        style.paragraph_format.line_spacing = 1.14


def configure_section(
    document: Document,
    header_text: str = "\t探海灵眸 Oceanus  |  项目计划书 v2.1",
) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.32)
    section.footer_distance = Inches(0.32)
    section.different_first_page_header_footer = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), 2)
    run = header.add_run(header_text)
    set_run_font(run, BODY_FONT, 8)
    run.font.color.rgb = RGBColor(90, 90, 90)

    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.65), 2)
    run = footer.add_run("证据边界：E0-E4")
    set_run_font(run, BODY_FONT, 8)
    run.font.color.rgb = RGBColor(90, 90, 90)
    run = footer.add_run("\t第 ")
    set_run_font(run, BODY_FONT, 8.5)
    run = footer.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, text, end])
    set_run_font(run, BODY_FONT, 8.5)
    run = footer.add_run(" 页")
    set_run_font(run, BODY_FONT, 8.5)


def add_inline(paragraph, text: str, base_size: float = 10.5) -> None:
    token_pattern = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    position = 0
    for match in token_pattern.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, BODY_FONT, base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, BODY_FONT, base_size, True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, MONO_FONT, max(base_size - 0.5, 8.0))
        else:
            link_match = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            run = paragraph.add_run(link_match.group(1) if link_match else token)
            set_run_font(run, BODY_FONT, base_size)
            if link_match:
                run.font.color.rgb = RGBColor(31, 78, 121)
                run.font.underline = True
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, BODY_FONT, base_size)


def add_cover(
    document: Document,
    metadata: dict[str, str],
    h1_count: int,
    cover_title: str | None = None,
    cover_subtitle: str = "区域海上环卫治理系统",
    cover_statement: str = "国奖候选工程基线",
    cover_meta_rows: list[tuple[str, str]] | None = None,
    cover_boundary: str = (
        "本文档严格区分代码/测试、合成验证、用户真实验证和商业成交证据；"
        "所有对外数字必须回到证据台账。"
    ),
) -> None:
    for _ in range(3):
        document.add_paragraph()

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_border(title)
    title_text = cover_title or metadata.get("title", "探海灵眸 Oceanus 项目计划书")
    run = title.add_run(title_text)
    set_run_font(run, BODY_FONT, 30, True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    run = subtitle.add_run(cover_subtitle)
    set_run_font(run, BODY_FONT, 16, True)

    statement = document.add_paragraph()
    statement.alignment = WD_ALIGN_PARAGRAPH.CENTER
    statement.paragraph_format.space_after = Pt(24)
    run = statement.add_run(cover_statement)
    set_run_font(run, BODY_FONT, 13, True)

    meta_rows = cover_meta_rows or [
        ("版本", metadata.get("版本", "2.1")),
        ("编制日期", metadata.get("编制日期", "2026-09-19")),
        ("参赛赛道", metadata.get("参赛赛道", "海洋科创 / 海上智能装备与无人系统")),
        ("文档定位", metadata.get("文档定位", "研发、产品、商业、答辩、验收和智能体分派的唯一总控基线")),
        ("当前评分", "80/100（内部红队；不是赛事官方评分）"),
        ("章节数量", f"{h1_count} 个一级章节"),
    ]
    meta_table = document.add_table(rows=len(meta_rows), cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    for row_index, (key, value) in enumerate(meta_rows):
        cells = meta_table.rows[row_index].cells
        cells[0].width = Inches(1.25)
        cells[1].width = Inches(5.35)
        cells[0].text = key
        cells[1].text = value
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            set_cell_border(cell, BORDER_COLOR)
        set_cell_shading(cells[0], "EAF1F7")
        for paragraph in cells[0].paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_run_font(run, BODY_FONT, 9, True)
        for paragraph in cells[1].paragraphs:
            for run in paragraph.runs:
                set_run_font(run, BODY_FONT, 9)

    document.add_paragraph()
    boundary = document.add_paragraph()
    boundary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = boundary.add_run(cover_boundary)
    set_run_font(run, BODY_FONT, 9.5)
    run.font.color.rgb = RGBColor(80, 80, 80)
    boundary.paragraph_format.space_before = Pt(18)

    document.add_page_break()


def add_toc(document: Document, headings: list[str], page_map: dict[str, int]) -> None:
    heading = document.add_heading("目录", level=1)
    heading.paragraph_format.page_break_before = False

    paragraph = document.add_paragraph()
    run = paragraph.add_run("页码以本文件最终渲染结果为准。")
    set_run_font(run, BODY_FONT, 9)
    run.font.color.rgb = RGBColor(90, 90, 90)

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    header = table.rows[0]
    header.cells[0].text = "章节"
    header.cells[1].text = "页码"
    set_repeat_table_header(header)
    for cell in header.cells:
        set_cell_shading(cell, HEADER_FILL)
        set_cell_border(cell, BORDER_COLOR)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if cell is header.cells[1] else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                set_run_font(run, BODY_FONT, 9, True)
                run.font.color.rgb = RGBColor(255, 255, 255)

    for row_index, text in enumerate(headings):
        row = table.add_row()
        row.cells[0].text = text
        row.cells[1].text = str(page_map.get(_normalize_heading(text), "-"))
        row.cells[0].width = Inches(6.10)
        row.cells[1].width = Inches(0.84)
        for cell in row.cells:
            set_cell_border(cell, BORDER_COLOR)
            set_cell_margins(cell, top=45, bottom=45)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index % 2 == 1:
                set_cell_shading(cell, ALT_FILL)
        row.cells[0].paragraphs[0].paragraph_format.space_after = Pt(0)
        row.cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row.cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in row.cells[0].paragraphs[0].runs:
            set_run_font(run, BODY_FONT, 9, False)
        for run in row.cells[1].paragraphs[0].runs:
            set_run_font(run, BODY_FONT, 9, False)

    document.add_page_break()


def calc_table_widths(rows: list[list[str]], total_width: float = 6.94) -> list[float]:
    columns = max(len(row) for row in rows)
    scores: list[float] = []
    for column in range(columns):
        values = [row[column] if column < len(row) else "" for row in rows]
        typical = sorted(len(value) for value in values)[max(0, len(values) // 2 - 1)]
        scores.append(max(4.0, min(float(typical), 42.0)))
    score_sum = sum(scores)
    widths = [total_width * score / score_sum for score in scores]
    minimum = 0.62 if columns >= 4 else 0.80
    for index, width in enumerate(widths):
        if width < minimum:
            deficit = minimum - width
            widths[index] = minimum
            donor_indexes = sorted(
                (i for i in range(len(widths)) if i != index and widths[i] - minimum > deficit),
                key=lambda i: widths[i],
                reverse=True,
            )
            for donor in donor_indexes:
                move = min(deficit, widths[donor] - minimum)
                widths[donor] -= move
                deficit -= move
                if deficit <= 0:
                    break
    scale = total_width / sum(widths)
    return [width * scale for width in widths]


def add_table(document: Document, rows: list[list[str]]) -> None:
    columns = max(len(row) for row in rows)
    table = document.add_table(rows=0, cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = calc_table_widths(rows)
    base_size = 8.2 if columns >= 4 else 8.8

    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        if row_index == 0:
            set_repeat_table_header(row)
        for column_index in range(columns):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.text = values[column_index] if column_index < len(values) else ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, top=38, start=70, bottom=38, end=70)
            set_cell_border(cell, BORDER_COLOR)
            if row_index == 0:
                set_cell_shading(cell, HEADER_FILL)
                alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                if row_index % 2 == 0:
                    set_cell_shading(cell, ALT_FILL)
                content = values[column_index] if column_index < len(values) else ""
                alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                    if len(content) <= 8 and not re.search(r"[。；：]", content)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignment
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.clear()
            add_inline(paragraph, values[column_index] if column_index < len(values) else "", base_size)
            for run in paragraph.runs:
                if row_index == 0:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
                else:
                    run.font.color.rgb = RGBColor(0, 0, 0)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_code_block(document: Document, lines: Iterable[str]) -> None:
    code_lines = list(lines)
    for index, line in enumerate(code_lines):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.10)
        paragraph.paragraph_format.right_indent = Inches(0.10)
        paragraph.paragraph_format.space_before = Pt(4 if index == 0 else 0)
        paragraph.paragraph_format.space_after = Pt(4 if index == len(code_lines) - 1 else 0)
        paragraph.paragraph_format.line_spacing = 1.0
        set_paragraph_shading(paragraph, CODE_FILL)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, MONO_FONT, 8.0)


def add_blocks(document: Document, blocks: list[Block]) -> None:
    first_h1 = True
    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                paragraph = document.add_heading(block.text, level=1)
                if not first_h1:
                    paragraph.paragraph_format.page_break_before = True
                first_h1 = False
            elif block.level == 2:
                document.add_heading(block.text, level=2)
            else:
                document.add_heading(block.text, level=3)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            add_inline(paragraph, block.text)
        elif block.kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.right_indent = Inches(0.18)
            paragraph.paragraph_format.space_before = Pt(3)
            paragraph.paragraph_format.space_after = Pt(7)
            add_inline(paragraph, block.text)
            for run in paragraph.runs:
                run.italic = True
        elif block.kind == "list":
            if block.ordered:
                # Keep the source list number instead of Word's document-wide
                # automatic sequence, which would turn independent lists into
                # one continuously numbered list.
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Inches(0.28 + 0.18 * block.level)
                paragraph.paragraph_format.first_line_indent = Inches(-0.22)
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.14
                marker = paragraph.add_run(f"{block.marker} ")
                set_run_font(marker, BODY_FONT, 10.5)
                add_inline(paragraph, block.text)
            else:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.left_indent = Inches(0.24 + 0.18 * block.level)
                paragraph.paragraph_format.first_line_indent = Inches(-0.18)
                add_inline(paragraph, block.text)
        elif block.kind == "table" and block.rows:
            add_table(document, block.rows)
        elif block.kind == "code" and block.lines is not None:
            add_code_block(document, block.lines)
        elif block.kind == "spacer":
            continue


def build_docx(
    source_path: Path,
    output_path: Path,
    heading_pages_path: Path | None = None,
    header_text: str = "\t探海灵眸 Oceanus  |  项目计划书 v2.1",
    cover_title: str | None = None,
    cover_subtitle: str = "区域海上环卫治理系统",
    cover_statement: str = "国奖候选工程基线",
    cover_meta_rows: list[tuple[str, str]] | None = None,
    cover_boundary: str = (
        "本文档严格区分代码/测试、合成验证、用户真实验证和商业成交证据；"
        "所有对外数字必须回到证据台账。"
    ),
    core_title: str = "探海灵眸 Oceanus 项目计划书",
    core_subject: str = "国奖候选工程基线与智能体开发分工方案",
    core_author: str = "Oceanus 项目组",
    core_comments: str = "严格区分 E0-E4 证据，不以规划替代真实现场与商业验证。",
) -> dict[str, object]:
    source = source_path.read_text(encoding="utf-8")
    metadata, blocks = parse_markdown(source)
    headings = [block.text for block in blocks if block.kind == "heading" and block.level == 1]
    page_map: dict[str, int] = {}
    if heading_pages_path and heading_pages_path.exists():
        raw_map = json.loads(heading_pages_path.read_text(encoding="utf-8"))
        page_map = {_normalize_heading(key): int(value) for key, value in raw_map.items()}

    document = Document()
    configure_styles(document)
    configure_section(document, header_text=header_text)
    add_cover(
        document,
        metadata,
        len(headings),
        cover_title=cover_title,
        cover_subtitle=cover_subtitle,
        cover_statement=cover_statement,
        cover_meta_rows=cover_meta_rows,
        cover_boundary=cover_boundary,
    )
    add_toc(document, headings, page_map)
    add_blocks(document, blocks)

    core = document.core_properties
    core.title = core_title
    core.subject = core_subject
    core.author = core_author
    core.comments = core_comments

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {
        "output": str(output_path),
        "headings": headings,
        "page_map_entries": len(page_map),
        "blocks": len(blocks),
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--heading-pages", type=Path, default=None)
    args = parser.parse_args()
    result = build_docx(args.source, args.output, args.heading_pages)
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
