# -*- coding: utf-8 -*-
r"""把校级「辅助材料」整理进连江提交包 → 连江县创新创业大赛\佐证材料\。

- 01/02/03/05 整目录复制（保留编号结构）
- 04_检测效果图：补入机械臂受控拾取实拍.mp4（来自 _backup_selection_20261010），
  识别样例图重排编号（03→04 等），使文件与 00_检测效果说明.docx 完全对应
- 文件名：scripts/complete_lianjiang_supporting.py
"""
import shutil
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

SRC = Path(r"C:\Users\Liii\Desktop\计算与信息科学学院（智能科学与技术）-探海灵眸-李涌翔-校级选拔材料\辅助材料")
BACKUP = Path(r"c:\Users\Liii\Desktop\seahawk\_backup_selection_20261010")
DST = Path(r"C:\Users\Liii\Desktop\连江县创新创业大赛\佐证材料")

DST.mkdir(exist_ok=True)

# ---------- 1) 整目录复制 ----------
for name in ["01_知识产权-软著", "02_用户访谈-连江县海渔局",
             "03_供应商询价调研-浙江蓝舶游艇科技公司", "05_评测与验证"]:
    src, dst = SRC / name, DST / name
    if dst.exists():
        shutil.rmtree(dst)
    shutil.copytree(src, dst)
    print("[COPY]", name)

# ---------- 2) 04_检测效果图：重排编号 + 补 mp4 ----------
d4 = DST / "04_检测效果图"
d4.mkdir(exist_ok=True)
s4 = SRC / "04_检测效果图"
pairs = [
    (s4 / "01_监测大屏截图.png", "01_监测大屏截图.png"),
    (s4 / "02_机械臂拾取作业.png", "02_机械臂拾取作业.png"),
    (BACKUP / "04_检测效果图" / "03_机械臂受控拾取实拍.mp4", "03_机械臂受控拾取实拍.mp4"),
    (s4 / "03_水面多目标检测.jpg", "04_水面多目标检测.jpg"),
    (s4 / "04_岸滩多目标检测.jpg", "05_岸滩多目标检测.jpg"),
    (s4 / "05_岸滩泡沫塑料检测.jpg", "06_岸滩泡沫塑料检测.jpg"),
]
for src, name in pairs:
    assert src.exists(), "缺源文件: %s" % src
    shutil.copy2(src, d4 / name)
    print("[COPY]", name)

# 说明文档复制（随后修正）
shutil.copy2(s4 / "00_检测效果说明.docx", d4 / "00_检测效果说明.docx")

# ---------- 3) 修正 00_检测效果说明.docx（连江副本） ----------
doc = Document(d4 / "00_检测效果说明.docx")


def rewrite_para(p, full: str) -> None:
    """保留首 run 格式整段重写"""
    assert p.runs, "空段落无法重写"
    p.runs[0].text = full
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


# 3.1 标题行
for p in doc.paragraphs:
    if "系统运行截图 3 项" in p.text:
        rewrite_para(p, p.text.replace("系统运行截图 3 项 + 识别效果样例 4 张",
                                       "系统运行记录 3 项 + 识别效果样例 3 张"))
        print("[FIX] 标题行已更新")

# 3.2 表1（识别效果样例）：删「水面塑料容器」行，06/07 → 05/06
t1 = doc.tables[1]


def set_cell(cell, text: str) -> None:
    p = cell.paragraphs[0]
    assert p.runs, "空单元格"
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


del_done = False
for row in t1.rows:
    c0 = row.cells[0].text.strip()
    c1 = row.cells[1].text.strip()
    if "05_水面塑料容器检测" in c1:
        row._element.getparent().remove(row._element)
        del_done = True
        print("[FIX] 删除行: %s" % c1)
    elif c1.startswith("06_"):
        set_cell(row.cells[0], "05")
        set_cell(row.cells[1], c1.replace("06_", "05_", 1))
        print("[FIX] 06→05:", c1)
    elif c1.startswith("07_"):
        set_cell(row.cells[0], "06")
        set_cell(row.cells[1], c1.replace("07_", "06_", 1))
        print("[FIX] 07→06:", c1)
assert del_done, "未找到待删行"

doc.save(d4 / "00_检测效果说明.docx")

# ---------- 4) 生成 00_佐证材料索引.docx ----------
def zh(run, name="宋体"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


idx = Document()
st = idx.styles["Normal"]
st.font.name = "宋体"
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

h = idx.add_paragraph()
r = h.add_run("探海灵眸 Oceanus · 佐证材料索引")
r.bold = True
r.font.size = Pt(16)
zh(r, "黑体")
h.alignment = 1  # 居中

h2 = idx.add_paragraph()
r = h2.add_run("连江县首届创新创业大赛（海洋科创赛道）· 2026 年 10 月 · 申报人：李涌翔")
r.font.size = Pt(9)
zh(r)
h2.alignment = 1

def body(text, bold=False, space=4):
    p = idx.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    zh(r)
    p.paragraph_format.space_after = Pt(space)
    return p

body("本目录为项目佐证材料，用于核验《报名表》《创业项目计划书》《路演 PPT》中引用的数据与证据；"
     "全部内容可溯源，逐项标注证据等级（E0 测算 / E1 代码与确定性仿真 / E2 多模块联调；更高等级证据单独登记）。")

body("一、01_知识产权-软著 —— 对应「3 项软件著作权登记申请中」", bold=True, space=2)
body("软著文档鉴别材料（软件说明书 22 页）+ 程序鉴别材料（源程序 60 页）。探海灵眸海洋环境治理智能体软件 V1.0"
     "（Oceanus）已完成鉴别材料制作并提交中国版权保护中心，尚未取得登记证书（聆心、BlindGuard 2 项同步申请中）。")

body("二、02_用户访谈-连江县海渔局 —— 对应「海渔局访谈完成」（E3）", bold=True, space=2)
body("访谈纪要 + 访谈过程截图 5 张。要点：①痛点证实（人工巡查发现、风飘垃圾盲区、县级无确切清运台账）；"
     "②需求认可（乡镇对「自动识别＋清理」明确表态认可）；③采购口径校准（现行全县域 13 乡镇「全年总包制」）。"
     "引用标注：连江县海渔局访谈（2026.10）。")

body("三、03_供应商询价调研-浙江蓝舶游艇科技公司 —— 对应「13 家 15 份询价」", bold=True, space=2)
body("询价汇总表 + 15 份询价原始记录（微信/平台截图与 PDF），其中询价 13、15 为供应商正式盖章报价单。"
     "覆盖 5 类核心部件：船体·浮台 / 推进器 / 电池组 / 边缘计算盒 / 摄像头。均为询价/测算口径（E0），非成交价。")

body("四、04_检测效果图 —— 对应感知链路运行效果", bold=True, space=2)
body("系统运行记录 3 项（监测大屏、机械臂拾取作业、机械臂受控拾取实拍视频）+ 识别效果样例 3 张"
     "（水面多目标、岸滩多目标、岸滩泡沫塑料）。识别结果为开源模型 YOLOv8s-World v2 推理输出，样例图为公开数据集场景。")

body("五、05_评测与验证 —— 对应「27/27 场景评测、129/129 集成验收、127 项故障注入、SC-FloatWaste 公开基准」", bold=True, space=2)
body("佐证说明 + 原始报告 JSON 7 份（含代码版本指纹、配置哈希与命令，可复跑、可复核）。"
     "①公开基准：监督微调通道 YOLO11 mAP50 0.514 / 零样本通道 YOLO-World 0.116（400 图 / 1,490 真值框）；"
     "②智能体固定场景评测 27/27（E1）；③真实状态集成验收 129/129（E2）；④受控故障注入自证 127 项（E1）；"
     "⑤自动化测试 1,061 项全部通过。")

body("口径声明", bold=True, space=2)
body("· 感知精度 not_evaluated，不虚报识别精度；42%～76% 仅指「时序链路误报抑制率」，不代表识别精度、召回率或准确率。\n"
     "· 公开基准评测为第三方公开数据集上的独立测试集精度，非连江现场验证、非本项目设备实测。\n"
     "· 商业数字均为测算/询价口径，不构成成交承诺；机械臂演示为 E2 受控实验，不宣称真实海域验证或已交付。")

idx.save(DST / "00_佐证材料索引.docx")
print("[OK] 索引已生成")

# ---------- 5) 打印最终目录树 ----------
print()
for f in sorted(DST.rglob("*")):
    depth = len(f.relative_to(DST).parts) - 1
    print("  " * depth + ("|-- " if depth else "") + f.name)