# -*- coding: utf-8 -*-
"""生成《情况说明（未设经营实体、无出资）》docx（连江大赛提交用）。"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(r"C:\Users\Liii\Desktop\连江县创新创业大赛\证明材料\情况说明_未设经营实体无出资_李涌翔.docx")


def para(doc, text, *, size=12, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, indent=True):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.first_line_indent = Pt(24)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def main() -> None:
    doc = Document()
    # 标题
    para(doc, "情 况 说 明", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER,
         space_after=16, indent=False)

    para(doc, "“数启海洋·智创未来”连江县首届创新创业大赛组委会：")
    para(doc,
         "本人李涌翔，男，身份证号 350825200511274116，系福州理工学院计算与信息科学学院"
         "智能科学与技术专业 2025 级在读本科生，学号 2501133021，为参赛项目"
         "“探海灵眸 Oceanus——海漂垃圾『感知—决策—执行』一体化智能治理系统”的申报人。")
    para(doc,
         "本项目为在校学生创新创业项目。截至本说明出具之日，本项目未设立任何经营实体"
         "（含企业、个体工商户等），不存在注册资本、实缴出资及股东出资情形，因此无法提供"
         "提交材料要求中的“申报人出资材料”（第 5 项），亦无需提交营业执照副本等经营实体"
         "资格证照（第 3 项）。")
    para(doc,
         "特此说明。以上情况真实无误，如有不实，本人愿承担相应责任。")
    para(doc, "联系电话：18259099729", space_after=24)
    para(doc, "申报人（签字）：＿＿＿＿＿＿＿＿＿＿", align=WD_ALIGN_PARAGRAPH.RIGHT,
         space_after=10, indent=False)
    para(doc, "日期：2026 年 10 月 ＿＿ 日", align=WD_ALIGN_PARAGRAPH.RIGHT, indent=False)

    doc.save(OUT)
    print("[OK]", OUT.name)


if __name__ == "__main__":
    main()