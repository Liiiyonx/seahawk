#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成「探海灵眸_章节页_浅色版.pptx」。

设计约定（与 chapters_light.html 一一对应）：
  · 每页 = 背景图（_bg/*.png：卡片/柔光/水印/品牌/基线，**不含正文**）
          + 可编辑文字层（文本框坐标/字号/字距均取自 HTML 的设计坐标）
  · 文字层在 WPS / PowerPoint 中可直接改（字体：微软雅黑 + Segoe UI）
  · 设计坐标系：1920×1080 px → 13.333×7.5 in（1px = 6350 EMU，精确 1:1）

依赖：python-pptx
运行：.venv-analysis\\Scripts\\python.exe 项目文档\\演示章节页\\浅色版\\make_pptx.py
"""
from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 6350  # 1 px = 6350 EMU（1920px = 13.333in = 12192000 EMU）
PAGE_W, PAGE_H = 1920, 1080
# 字号换算：1920px 映射 13.333in ⇒ 144 dpi ⇒ 1 设计px = 0.5pt
# （注意：不是 96dpi 的 0.75 —— 用错会让文字放大 1.5 倍）
K = 0.5

INK = "1C1C1E"   # --text-main
SUB = "636366"   # --text-sub
DIM = "8E8E93"   # --text-dim
BLUE = "007AFF"  # --c-primary（平台浅色模式主色）

CJK = "Microsoft YaHei"
LATIN = "Segoe UI"

BASE = Path(__file__).resolve().parent
OUT = BASE / "探海灵眸_章节页_浅色版.pptx"


def emu(px: float) -> Emu:
    return Emu(int(round(px * PX)))


def style_run(run, *, size, color, bold=False, spc=None):
    """size/spc 传「设计像素」；内部换算成 pt（px × 0.5）。"""
    f = run.font
    f.size = Pt(size * K)
    f.bold = bold
    f.color.rgb = RGBColor.from_string(color)
    f.name = LATIN
    rpr = run._r.get_or_add_rPr()
    # 中文字体走 a:ea（紧跟 a:latin，符合 schema 顺序）
    latin = rpr.find(qn("a:latin"))
    ea = rpr.find(qn("a:ea"))
    if ea is None:
        ea = rpr.makeelement(qn("a:ea"), {})
        if latin is not None:
            latin.addnext(ea)
        else:
            rpr.append(ea)
    ea.set("typeface", CJK)
    if spc is not None:
        rpr.set("spc", str(int(round(spc * K * 100))))


def add_text(slide, *, x, y, w, runs, align="l", box_h=80):
    """加一个单行文本框；runs = [(text, dict(size=, color=, bold=, spc=)), ...]。"""
    box = slide.shapes.add_textbox(emu(x), emu(y), emu(w), emu(box_h))
    tf = box.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT if align == "l" else PP_ALIGN.CENTER
    p.line_spacing = 1.0
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    for text, opts in runs:
        r = p.add_run()
        r.text = text
        style_run(r, **opts)
    return box


def R(text, size, color, *, bold=False, spc=0.0):
    return (text, {"size": size, "color": color, "bold": bold, "spc": spc})


# 章节内容（与 chapters_light.html 完全一致）
CHAPTERS = [
    dict(num="01", en="· THE PROBLEM",
         h1="近岸垃圾之困",
         sub="发现靠人、处置靠等 —— 传统治理方式跟不上垃圾涌入的速度"),
    dict(num="02", en="· THE PLATFORM",
         h1="一屏观全局",
         sub="四路实时视频 + 3D 海域态势 —— 事件、工单、设备、报表尽在一屏"),
    dict(num="03", en="· HUMAN IN THE LOOP",
         h1="人机协同 · 从工单到执行",
         sub="工单全生命周期与执行仿真全链路 —— 智能助手与知识智能体随问随答"),
    dict(num="04", en="· SAFETY GUARDRAIL",
         h1="安全底线 · 审批与拦截",
         sub="高风险写入必须人工批准，证据不足主动拦截 —— 批准之前，智能体一步都不会动"),
    dict(num="05", en="· PLUGGABLE EXECUTION",
         h1="软硬闭环 · 从决策到执行",
         sub="同一套任务与消息契约 —— 无人机、水面机器人、机械臂皆可插拔"),
]


def build():
    prs = Presentation()
    prs.slide_width = emu(PAGE_W)
    prs.slide_height = emu(PAGE_H)
    blank = prs.slide_layouts[6]

    def new_slide(bg_name: str):
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(
            str(BASE / "_bg" / bg_name), 0, 0,
            width=emu(PAGE_W), height=emu(PAGE_H),
        )
        return slide

    # ---------- 0 片名页 ----------
    s = new_slide("title.png")
    add_text(s, x=0, y=316, w=PAGE_W, align="c",
             runs=[R("MULTI-AGENT COASTAL WASTE GOVERNANCE", 21, DIM, spc=11)])
    add_text(s, x=0, y=368, w=PAGE_W, align="c",
             runs=[R("探海灵眸", 170, INK, bold=True, spc=22)])
    add_text(s, x=0, y=586, w=PAGE_W, align="c",
             runs=[R("O C E A N U S", 46, BLUE, spc=32)])
    add_text(s, x=0, y=722, w=PAGE_W, align="c",
             runs=[R("海漂垃圾全链路智能治理平台 ", 32, INK, spc=5),
                   R("·", 32, BLUE, spc=5),
                   R(" 平台演示", 32, DIM, spc=5)])

    # ---------- 1~5 章节页 ----------
    for i, ch in enumerate(CHAPTERS, start=1):
        s = new_slide(f"{i}.png")
        add_text(s, x=202, y=600, w=1300,
                 runs=[R(f"CHAPTER {ch['num']} ", 22, BLUE, bold=True, spc=6),
                       R(ch["en"], 22, DIM, spc=4)])
        add_text(s, x=140, y=660, w=1700, box_h=140,
                 runs=[R(ch["h1"], 105, INK, bold=True, spc=2)])
        add_text(s, x=140, y=812, w=1700,
                 runs=[R(ch["sub"], 30, SUB, spc=1)])

    # ---------- 6 结尾标版 ----------
    s = new_slide("outro.png")
    add_text(s, x=0, y=368, w=PAGE_W, align="c",
             runs=[R("OCEANUS · MULTI-AGENT PLATFORM", 21, DIM, spc=11)])
    add_text(s, x=0, y=430, w=PAGE_W, align="c",
             runs=[R("有据可依 ", 46, INK, bold=True, spc=8),
                   R("·", 46, BLUE, bold=True, spc=8),
                   R(" 有人把关 ", 46, INK, bold=True, spc=8),
                   R("·", 46, BLUE, bold=True, spc=8),
                   R(" 有迹可循", 46, INK, bold=True, spc=8)])
    add_text(s, x=0, y=588, w=PAGE_W, align="c",
             runs=[R("让每一次近岸垃圾处置，都有完整的智能协同闭环", 30, SUB, spc=3)])
    add_text(s, x=0, y=678, w=PAGE_W, align="c",
             runs=[R("探海灵眸 · 2026 平台演示", 20, DIM, spc=6)])

    prs.save(str(OUT))
    print(f"[OK] {OUT.name}（{len(prs.slides._sldIdLst)} 页）")


if __name__ == "__main__":
    build()